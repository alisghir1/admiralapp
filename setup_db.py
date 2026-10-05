import sqlite3
import openpyxl
import os

DB_FILE = 'suivi_atelier.db'
EXCEL_FILE = r'C:\Users\Ali\Downloads\BASE SAGE ADMIRAL.xlsx'

def setup_database():
    print("Initialisation de la nouvelle base de données MES...")
    
    # Connect to the database (creates it if it doesn't exist)
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    
    # 1. Drop existing tables to start fresh
    tables = ['Produits', 'Operations', 'Postes', 'Gammes', 'CdeDetail', 'Scans', 'OF_Gammes']
    for table in tables:
        cursor.execute(f"DROP TABLE IF EXISTS {table}")
        
    # 2. Create the new schema
    cursor.execute('''
        CREATE TABLE Produits (
            RefProd TEXT PRIMARY KEY,
            DesProd TEXT
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE Operations (
            RefOp TEXT PRIMARY KEY,
            LibelOp TEXT
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE Postes (
            CodPost TEXT PRIMARY KEY,
            DesPost TEXT,
            RefOp TEXT,
            FOREIGN KEY(RefOp) REFERENCES Operations(RefOp)
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE Gammes (
            ID INTEGER PRIMARY KEY AUTOINCREMENT,
            RefProd TEXT,
            CodeOp TEXT,
            Ordre INTEGER,
            FOREIGN KEY(RefProd) REFERENCES Produits(RefProd),
            FOREIGN KEY(CodeOp) REFERENCES Operations(RefOp)
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE CdeDetail (
            ID_OF TEXT PRIMARY KEY,
            CDE TEXT,
            PROD TEXT,
            DESIGN TEXT,
            OBSERV TEXT,
            Qte INTEGER DEFAULT 1,
            FOREIGN KEY(PROD) REFERENCES Produits(RefProd)
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE OF_Gammes (
            ID_OF TEXT,
            Ordre INTEGER,
            CodeOp TEXT,
            Poste_Assigne TEXT,
            PRIMARY KEY (ID_OF, Ordre),
            FOREIGN KEY(ID_OF) REFERENCES CdeDetail(ID_OF),
            FOREIGN KEY(CodeOp) REFERENCES Operations(RefOp)
        )
    ''')
    
    # We redefine Scans to link to the ID_OF
    cursor.execute('''
        CREATE TABLE Scans (
            ID_Scan INTEGER PRIMARY KEY AUTOINCREMENT,
            ID_OF TEXT NOT NULL,
            Poste TEXT NOT NULL,
            Timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(ID_OF) REFERENCES CdeDetail(ID_OF),
            FOREIGN KEY(Poste) REFERENCES Postes(CodPost)
        )
    ''')
    
    print("Schéma créé avec succès.")

    # 3. Import Data from Excel
    if os.path.exists(EXCEL_FILE):
        print(f"Lecture du fichier Excel : {EXCEL_FILE}")
        wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
        
        # Import Produits
        if 'PRODUIT' in wb.sheetnames:
            ws = wb['PRODUIT']
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or not row[0]: continue # Skip header or empty rows
                cursor.execute("INSERT OR IGNORE INTO Produits (RefProd, DesProd) VALUES (?, ?)", (str(row[0]), str(row[1]) if row[1] else ""))
            print("Produits importés.")

        # Import Operations
        if 'OPERATION' in wb.sheetnames:
            ws = wb['OPERATION']
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or not row[0]: continue
                cursor.execute("INSERT OR IGNORE INTO Operations (RefOp, LibelOp) VALUES (?, ?)", (str(row[0]), str(row[1]) if row[1] else ""))
            print("Opérations importées.")

        # Import Postes
        if 'POSTE' in wb.sheetnames:
            ws = wb['POSTE']
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or not row[0]: continue
                codpost = str(row[0]).strip()
                despost = str(row[1]).strip() if row[1] else ""
                # Infer RefOp from CodPost (e.g. DEC1 -> DEC)
                # This is a heuristic since RefOp isn't explicitly in the POSTE sheet
                refop = ''.join([c for c in codpost if not c.isdigit()]).strip()
                cursor.execute("INSERT OR IGNORE INTO Postes (CodPost, DesPost, RefOp) VALUES (?, ?, ?)", (codpost, despost, refop))
            print("Postes importés.")

        # Import Gammes (GAMPROD)
        if 'GAMPROD' in wb.sheetnames:
            ws = wb['GAMPROD']
            current_prod = None
            ordre = 1
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0: continue
                
                col_refprod = row[0]
                col_codop = row[2]
                
                if col_refprod: # New product routing starts
                    current_prod = str(col_refprod).strip()
                    ordre = 1
                elif col_codop and current_prod: # Operation for current product
                    cursor.execute("INSERT INTO Gammes (RefProd, CodeOp, Ordre) VALUES (?, ?, ?)", (current_prod, str(col_codop).strip(), ordre))
                    ordre += 1
            print("Gammes importées.")

        # Import CDEDETAIL (Initial OF creation)
        if 'CDEDETAIL' in wb.sheetnames:
            ws = wb['CDEDETAIL']
            of_counter = 1000
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or not row[0]: continue
                cde = str(row[0]).strip()
                prod = str(row[1]).strip() if row[1] else ""
                design = str(row[2]).strip() if row[2] else ""
                observ = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                
                id_of = f"OF-{of_counter}"
                of_counter += 1
                qte = 1 # Par défaut, car pas dans le Excel d'origine
                
                cursor.execute("INSERT INTO CdeDetail (ID_OF, CDE, PROD, DESIGN, OBSERV, Qte) VALUES (?, ?, ?, ?, ?, ?)", (id_of, cde, prod, design, observ, qte))
                
                # Générer la gamme pour cet OF
                cursor.execute("SELECT CodeOp, Ordre FROM Gammes WHERE RefProd = ? ORDER BY Ordre", (prod,))
                gamme_standard = cursor.fetchall()
                for step in gamme_standard:
                    cursor.execute("INSERT INTO OF_Gammes (ID_OF, Ordre, CodeOp) VALUES (?, ?, ?)", (id_of, step[1], step[0]))
                    
            print("Commandes Détails (OF) et Gammes OF importées.")

    conn.commit()
    conn.close()
    print("Base de données mise à jour avec succès.")

if __name__ == "__main__":
    setup_database()
